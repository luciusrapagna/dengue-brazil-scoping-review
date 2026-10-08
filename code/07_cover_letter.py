"""Cover letter to the Editors-in-Chief (journal recommends explaining novelty)."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_COLOR_INDEX
import re

ROOT = Path(__file__).resolve().parents[1]
TEXT = [
    "[[INSERT: city]], [[INSERT: date]]",
    "To the Editors-in-Chief, Ciência & Saúde Coletiva",
    "Dear Editors,",
    "We submit the review article \"Four decades of dengue epidemics in Brazil: a scoping review of serotype succession, severity and evidence gaps\" for consideration in Ciência & Saúde Coletiva.",
    "Brazil recorded the largest dengue epidemic in its history in 2024. Previous reviews covered specific periods, regions or outcomes. To our knowledge, this is the first review to integrate 45 years of evidence on dengue epidemics in Brazil, published in English, Portuguese and Spanish, and to analyse the literature itself as an object of study. The article brings three contributions to collective health: (1) a periodization of Brazilian dengue epidemics into five phases organized by serotype succession, tested against an original serotype-year evidence map; (2) evidence that knowledge production is concentrated in the Southeast and does not follow the distribution of burden; and (3) evidence that the reporting of serotypes in titles and abstracts decreased significantly over time, with a proposal of minimum reporting items that is directly relevant to vaccination, Wolbachia deployment and epidemic preparedness.",
    "The analysis is anchored in the immunological framework of serotype succession and in the socio-spatial notion of the pathogenic complex, and places the findings in dialogue with national and international literature. All data and code are openly available ([[INSERT: Zenodo DOI]]), in line with the journal's open science policy.",
    "[[INSERT: statement on the previous submission, if appropriate: this manuscript is a substantially new study, with a new design, scope and analysis, compared with a previously submitted narrative review on dengue in the state of Rio de Janeiro.]]",
    "In accordance with the journal's policy, the use of artificial intelligence is described at the end of the Methods section. The manuscript is not under consideration elsewhere, and all authors approved the submitted version and declare [[INSERT: no conflicts of interest]].",
    "Sincerely,",
    "[[INSERT: corresponding author, on behalf of all authors]]",
]

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
for t in TEXT:
    p = doc.add_paragraph()
    for part in re.split(r"(\[\[[^\]]+\]\])", t):
        if part.startswith("[["):
            r = p.add_run("[" + part[2:-2] + "]"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part:
            p.add_run(part)
doc.save(ROOT / "manuscript" / "Cover_letter.docx")
print("Cover letter written")
