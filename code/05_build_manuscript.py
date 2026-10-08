"""Build the submission files for Ciência & Saúde Coletiva.

Reads manuscript/manuscript_source.md (citation keys such as [@key1;@key2]
and placeholders [[INSERT ...]]) and writes:
  manuscript/Manuscript_blinded.docx   - main text, references, tables, figures
  manuscript/Title_page.docx           - identification of authors (placeholders)
  tables/Table1_chronology.docx, tables/Table2_characteristics.docx (editable)
  manuscript/reference_list.txt        - numbered reference list (audit)

Journal rules applied: Times New Roman 12, double spacing, 2.5 cm margins,
unnumbered headings, Vancouver references numbered by order of citation and
cited as superscript numbers, all authors listed (et al. only above 25 authors),
NLM journal abbreviations; abstract <= 1,400 characters; text from
"Introduction" to the last reference <= 45,000 characters (review article).
"""
from pathlib import Path
import json
import re
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "manuscript" / "manuscript_source.md"
META = ROOT / "data" / "processed" / "references_metadata.json"
OUTM = ROOT / "manuscript"
OUTT = ROOT / "tables"
FIG = ROOT / "figures"

ORIGINAL_TITLES = {  # titles in the original language of publication (journal rule)
    "osanai1983": "Surto de dengue em Boa Vista, Roraima. Nota prévia",
    "vasconcelos1998": "Epidemia de dengue em Fortaleza, Ceará: inquérito soro-epidemiológico aleatório",
    "casali2004": "A epidemia de dengue/dengue hemorrágico no município do Rio de Janeiro, 2001/2002",
    "passos2004": ("Diferenças clínicas observadas em pacientes com dengue causadas por diferentes sorotipos "
                   "na epidemia de 2001/2002, ocorrida no município do Rio de Janeiro"),
    "toledo2006": "Confiabilidade do diagnóstico final de dengue na epidemia 2001-2002 no Município do Rio de Janeiro, Brasil",
    "lima2007": ("Dengue: inquérito populacional para pesquisa de anticorpos e vigilância virológica "
                 "no Município de Campinas, São Paulo, Brasil"),
    "heinen2015": "Dengue outbreak in Mato Grosso State, midwestern Brazil",
    "oliveiraMA2018": ("El papel de los flujos interregionales en la diseminación de epidemias de dengue "
                       "en una ciudad de clima tropical"),
    "montenegro2006": "Aspectos clínicos e epidemiológicos da epidemia de dengue no Recife, PE, em 2002",
    "melo2010": "Progressão da circulação do vírus do dengue no Estado da Bahia, 1994-2000",
    "marques2020": ("Avaliação da não completude das notificações compulsórias de dengue registradas por município "
                    "de pequeno porte no Brasil"),
    "moraes2009": ("Análise da concordância dos dados de mortalidade por dengue em dois sistemas nacionais de "
                   "informação em saúde, Brasil, 2000-2005"),
    "costa2011": ("Dengue: aspectos epidemiológicos e o primeiro surto ocorrido na região do Médio Solimões, "
                  "Coari, Estado do Amazonas, no período de 2008 a 2009"),
}
MANUAL_REFS = {
    "catao2012": "Catão RC. Dengue no Brasil: abordagem geográfica na escala nacional. São Paulo: Cultura Acadêmica; 2012.",
}
JOURNAL_FIX = {"International Journal of Social Research Methodology": "Int J Soc Res Methodol"}
AUTHOR_FIX = {"Teixeira Mda G": "Teixeira MG"}

# Table 1: chronology of epidemic phases (synthesis of the included studies)
TABLE1 = [
    ("I. Re-emergence\n1981–1989", "DENV-1 and DENV-4 (Boa Vista, 1981–1982); DENV-1 (Rio de Janeiro, 1986)",
     "Localized explosive epidemics in immunologically naïve urban populations; classic dengue; "
     "limited accuracy of clinical case definitions",
     "[@osanai1983;@bezerra2021;@schatzmayr1986;@miagostovich1993;@dietz1990]"),
    ("II. Co-circulation\n1990–2000", "DENV-2 introduced (Rio de Janeiro, 1990); DENV-1/DENV-2 co-circulation",
     "Spread to the Northeast and Centre; first dengue haemorrhagic fever cases; high seroprevalence and "
     "'silent' transmission; shift to nationwide endemic-epidemic circulation from 1994",
     "[@nogueira1990;@nogueira1999;@nogueira1995;@vasconcelos1995;@vasconcelos1998;@cunha1999;@teixeira2002;@siqueira2005]"),
    ("III. Severity\n2001–2009", "DENV-3 introduced (2000–2001); DENV-2 new lineage (2007–2008)",
     "Largest epidemics until then (Rio de Janeiro 2002: 288,245 cases); higher odds of shock with DENV-3; "
     "dispersal along road networks; shift of severe disease to children in 2008",
     "[@nogueira2001;@nogueira2005;@passos2004;@melo2007;@montenegro2006;@cordeiro2007;@nunes2016;@teixeira2013]"),
    ("IV. Hyperendemicity\n2010–2019", "DENV-4 re-introduced (Roraima, 2010); four serotypes co-circulating",
     "DENV-4 waves in 2012–2013; yearly replacement of dominant serotype and lineages; large unreported "
     "fraction; rising burden and deaths; record local incidence in the South",
     "[@temporao2011;@ramalho2018;@heinen2015;@faria2017;@oliveira2018;@jesus2020;@araujo2017;@andrioli2020]"),
    ("V. Record epidemics and new tools\n2020–2026", "DENV-1 and DENV-2 predominant; DENV-3 re-emergence (lineage 3III_B.3.2, 2023)",
     "Expansion to the South; 2024 epidemic with about 6.6 million probable cases; vaccine and Wolbachia "
     "effectiveness evidence; delayed emergency responses",
     "[@silva2025;@gularte2023;@gurgel2024;@ferreira2026;@ranzani2025;@anders2025;@barberia2026]"),
]
FIG_LEGENDS = [
    ("Figure1_serotype_evidence_map", "Figure 1. Serotype-year evidence map of dengue epidemics in Brazil, 1981–2025.",
     "(A) Number of included studies with dated primary evidence for each epidemic year; shading indicates the five "
     "epidemic phases. (B) Number of studies reporting the circulation of each serotype by year; triangles indicate "
     "documented introductions or re-emergence. The most recent epidemic year documented by the included studies was 2025. Source: The authors, based on 106 included studies (495 serotype-year records)."),
    ("Figure2_geography_and_design", "Figure 2. Geographic distribution and design of the 286 included studies.",
     "(A) Number of studies by federative unit, grouped by macro-region; studies covering more than one unit are counted "
     "in each; pseudo-categories group studies whose federative unit was not specified. (B) Study design or focus by publication period. Source: The authors."),
    ("Figure3_serotype_reporting", "Figure 3. Source of serotype information in studies of dengue epidemics in Brazil, by publication period.",
     "Percentages refer to the 276 studies for which serotype information was applicable. Source: The authors."),
]


# ---------------------------------------------------------------------- refs
def expand_pages(p):
    if not p:
        return ""
    m = re.fullmatch(r"(\d+)-(\d+)", p)
    if m and len(m.group(2)) < len(m.group(1)):
        a, b = m.groups()
        return f"{a}-{a[:len(a) - len(b)]}{b}"
    return p


def format_ref(key, r):
    if key in MANUAL_REFS:
        return MANUAL_REFS[key]
    au = [AUTHOR_FIX.get(a, a) for a in r["authors"]]
    authors = ", ".join(au[:6]) + ", et al" if len(au) > 25 else ", ".join(au)
    title = ORIGINAL_TITLES.get(key, r["title"]).strip().rstrip(".")
    title = re.sub(r"^\[(.*)\]$", r"\1", title)
    j = JOURNAL_FIX.get(r["journal"], r["journal"])
    vol = r.get("volume") or ""
    iss = f"({r['issue']})" if r.get("issue") else ""
    pages = expand_pages(r.get("pages") or "")
    m = re.search(r"Suppl\s*0?(\d+)", f"{vol} {r.get('issue') or ''}")
    if m:  # journal style for supplements: 25(Supl. 1)
        vol, iss = re.match(r"\d+", vol).group(0), f"(Supl. {m.group(1)})"
    if re.fullmatch(r"S\d+-\d+", pages):
        a, b = pages[1:].split("-"); pages = f"S{a}-S{expand_pages(f'{a}-{b}').split('-')[1]}"
    tail = f"{vol}{iss}:{pages}" if pages else f"{vol}{iss}"
    end = "" if title.endswith("?") else "."
    return f"{authors}. {title}{end} {j} {r['year']}; {tail}."


class Citer:
    def __init__(self):
        self.order = []

    def numbers(self, keys):
        for k in keys:
            if k not in self.order:
                self.order.append(k)
        nums = sorted(self.order.index(k) + 1 for k in keys)
        out, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            out.append(f"{nums[i]}-{nums[j]}" if j - i >= 2 else ",".join(map(str, nums[i:j + 1])))
            i = j + 1
        return ",".join(out)


# ---------------------------------------------------------------------- docx
TOKEN = re.compile(r"(\[@[^\]]+\]|\[\[[^\]]+\]\]|\*[^*]+\*)")


def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE; pf.space_after = Pt(0); pf.space_before = Pt(0)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.5)
    add_page_numbers(doc.sections[0])


def add_page_numbers(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text
        r._r.append(e)


def add_rich(par, text, citer, size=None, sup=True):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("[@"):
            keys = [k.strip().lstrip("@") for k in part[1:-1].split(";")]
            run = par.add_run(citer.numbers(keys)); run.font.superscript = sup
        elif part.startswith("[["):
            run = par.add_run("[" + part[2:-2] + "]"); run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith("*") and part.endswith("*"):
            run = par.add_run(part[1:-1]); run.italic = True
        else:
            run = par.add_run(part)
        if size:
            run.font.size = Pt(size)
    return par


def heading(doc, text, level):
    p = doc.add_paragraph(); r = p.add_run(text if level == 2 else text.upper()); r.bold = True
    if level == 2:
        r.italic = True
    return p


def parse_source():
    meta, body = {}, []
    for line in SRC.read_text(encoding="utf-8").splitlines():
        if line.startswith("%"):
            k, v = line[1:].split(" ", 1); meta[k] = v.strip()
        elif line.strip():
            body.append(line.rstrip())
    return meta, body


def plain(text):
    """Text as it will appear (for character counting)."""
    return re.sub(r"\[@[^\]]+\]", "00", text).replace("[[", "[").replace("]]", "]").replace("*", "")


def write_table(doc, header, rows, citer, widths, font=9):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_rich(p, str(v), citer, size=font, sup=False)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return t


def main():
    meta, body = parse_source()
    refs_meta = json.load(open(META, encoding="utf-8"))
    citer = Citer()
    doc = Document(); setup(doc)

    # title, abstracts ---------------------------------------------------
    for k in ("TITLE", "TITLE_PT", "TITLE_ES"):
        assert len(meta[k]) <= 120, f"{k} has {len(meta[k])} characters (max 120)"
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(meta[k]); r.bold = (k == "TITLE")
    abstract_counts = {}
    for lab, ka, kk, kw in (("Abstract", "ABSTRACT", "KEYWORDS", "Keywords"),
                            ("Resumo", "ABSTRACT_PT", "KEYWORDS_PT", "Palavras-chave"),
                            ("Resumen", "ABSTRACT_ES", "KEYWORDS_ES", "Palabras clave")):
        doc.add_paragraph()
        p = doc.add_paragraph(); r = p.add_run(lab + " "); r.bold = True; p.add_run(meta[ka])
        p = doc.add_paragraph(); r = p.add_run(kw + " "); r.bold = True; p.add_run(meta[kk])
        n = len(f"{lab} {meta[ka]} {kw} {meta[kk]}")
        abstract_counts[lab] = n
        assert n <= 1400, f"{lab}: {n} characters (max 1,400)"
    doc.add_page_break()

    # main text ------------------------------------------------------------
    counted = []
    for line in body:
        if line.startswith("# "):
            heading(doc, line[2:], 1); counted.append(line[2:])
        elif line.startswith("## "):
            heading(doc, line[3:], 2); counted.append(line[3:])
        else:
            p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(1.25)
            add_rich(p, line, citer); counted.append(plain(line))
    # Table and figure citations in text are already in the body; tables are cited after.
    for k, lab in (("ACK", "Acknowledgements"), ("FUNDING", "Funding"), ("CONTRIB", "Authors' contributions"),
                   ("DATA", "Data availability")):
        heading(doc, lab, 2); p = doc.add_paragraph(); add_rich(p, meta[k], citer); counted += [lab, plain(meta[k])]

    # make table-only citations numbered after text citations
    for row in TABLE1:
        citer.numbers([k.lstrip("@") for k in re.findall(r"@(\w+)", row[3])])

    heading(doc, "References", 1); counted.append("References")
    ref_lines = []
    for i, k in enumerate(citer.order, 1):
        line = f"{i}. {format_ref(k, refs_meta.get(k, {}))}"
        ref_lines.append(line)
        p = doc.add_paragraph(line); p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
    counted += ref_lines
    (OUTM / "reference_list.txt").write_text("\n".join(ref_lines), encoding="utf-8")

    # tables -----------------------------------------------------------------
    t1_header = ["Phase and period", "Serotypes", "Main epidemiological features", "Key evidence"]
    t1_rows = [(a, b, c, d) for a, b, c, d in TABLE1]
    t2 = pd.read_csv(OUTT / "table2_characteristics.csv")
    tables = [
        ("Table 1. Phases of dengue epidemics in Brazil according to serotype succession, 1981–2026.",
         t1_header, t1_rows, [3.0, 3.6, 5.6, 2.8], "Source: The authors, based on the included studies (reference numbers in the last column)."),
        ("Table 2. Characteristics of the 286 studies on dengue epidemics in Brazil included in the review.",
         ["Characteristic", "n (%)"], list(t2.itertuples(index=False, name=None)), [11.0, 4.0],
         "Source: The authors. Percentages are calculated over the 286 studies; categories marked as not mutually exclusive may add up to more than 100%."),
    ]
    for title, header, rows, widths, source in tables:
        doc.add_page_break()
        p = doc.add_paragraph(); r = p.add_run(title); r.bold = True
        write_table(doc, header, [tuple("" if pd.isna(x) else x for x in row) for row in rows], citer, widths)
        p = doc.add_paragraph(); r = p.add_run(source); r.font.size = Pt(9)
        # standalone editable table file
        tdoc = Document(); setup(tdoc)
        p = tdoc.add_paragraph(); r = p.add_run(title); r.bold = True
        tc = Citer(); tc.order = list(citer.order)
        write_table(tdoc, header, [tuple("" if pd.isna(x) else x for x in row) for row in rows], tc, widths)
        p = tdoc.add_paragraph(); r = p.add_run(source); r.font.size = Pt(9)
        tdoc.save(OUTT / (title.split(".")[0].replace(" ", "") + ("_chronology" if "1" in title.split(".")[0] else "_characteristics") + ".docx"))

    # figures ---------------------------------------------------------------
    for fname, title, legend in FIG_LEGENDS:
        doc.add_page_break()
        p = doc.add_paragraph(); r = p.add_run(title); r.bold = True
        doc.add_picture(str(FIG / f"{fname}.png"), width=Cm(15))
        p = doc.add_paragraph(); r = p.add_run(legend); r.font.size = Pt(10)

    doc.save(OUTM / "Manuscript_blinded.docx")

    # title page -------------------------------------------------------------
    tp = Document(); setup(tp)
    for k in ("TITLE", "TITLE_PT", "TITLE_ES"):
        p = tp.add_paragraph(); r = p.add_run(meta[k]); r.bold = (k == "TITLE")
    items = [
        "Running title: [[INSERT: short title, e.g., Dengue epidemics and serotype succession in Brazil]]",
        "Authors: [[INSERT: full names of all authors (maximum of eight), in order]]",
        "Affiliations: [[INSERT: institution, department, city, state, country for each author]]",
        "ORCID: [[INSERT: ORCID iD of each author]]",
        "Corresponding author: [[INSERT: name, postal address and e-mail]]",
        "Authors' contributions: [[INSERT]]",
        "Funding: [[INSERT; CAPES Finance Code 001 if applicable]]",
        "Conflicts of interest: [[INSERT: e.g., The authors declare no conflicts of interest.]]",
        "Preprint: Not applicable. The manuscript has not been deposited in a preprint server.",
        "Data availability: The data set, codebook, search and screening logs and Python scripts that reproduce all figures, tables and results are openly available at https://github.com/luciusrapagna/dengue-brazil-scoping-review and archived in Zenodo (https://doi.org/10.5281/zenodo.23242684) under CC BY 4.0 (data) and MIT (code) licences.",
        f"Article type: Review article. Characters (Introduction to last reference, with spaces): {len(chr(10).join(counted)):,}",
    ]
    for it in items:
        p = tp.add_paragraph(); add_rich(p, it, Citer())
    tp.save(OUTM / "Title_page.docx")

    n_chars = len("\n".join(counted))
    report = dict(characters_intro_to_last_reference=n_chars, limit=45000, abstracts=abstract_counts,
                  references=len(citer.order), titles={k: len(meta[k]) for k in ("TITLE", "TITLE_PT", "TITLE_ES")})
    (OUTM / "compliance_report.json").write_text(json.dumps(report, indent=1), encoding="utf-8")
    print(json.dumps(report, indent=1))
    if n_chars > 45000:
        print("WARNING: text exceeds 45,000 characters")


if __name__ == "__main__":
    main()
