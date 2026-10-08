"""Supplementary material.

S1  Included studies with coded variables (xlsx)
S4  LILACS screening log (written by 04_prisma.py)
S2  Search log: sources, dates, strategies and records (xlsx)
S3  PRISMA-ScR checklist with placeholders for page numbers (docx)
"""
from pathlib import Path
import re
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_COLOR_INDEX

ROOT = Path(__file__).resolve().parents[1]
SUP = ROOT / "supplementary"
SUP.mkdir(exist_ok=True)

PRISMA_SCR = [
    ("TITLE", "1", "Title", "Identify the report as a scoping review."),
    ("ABSTRACT", "2", "Structured summary", "Background, objectives, eligibility criteria, sources of evidence, charting methods, results and conclusions."),
    ("INTRODUCTION", "3", "Rationale", "Rationale for the review in the context of what is already known."),
    ("INTRODUCTION", "4", "Objectives", "Questions and objectives addressed, with reference to key elements (population, concept, context)."),
    ("METHODS", "5", "Protocol and registration", "Whether a protocol exists, where it can be accessed and registration information."),
    ("METHODS", "6", "Eligibility criteria", "Characteristics of the sources of evidence used as eligibility criteria, with rationale."),
    ("METHODS", "7", "Information sources", "All information sources and the date each was last searched."),
    ("METHODS", "8", "Search", "Full electronic search strategy for at least one database."),
    ("METHODS", "9", "Selection of sources of evidence", "Process for selecting sources of evidence (screening and eligibility)."),
    ("METHODS", "10", "Data charting process", "Methods of charting data, whether forms were calibrated, independent or in duplicate."),
    ("METHODS", "11", "Data items", "All variables for which data were sought and any assumptions made."),
    ("METHODS", "12", "Critical appraisal (optional)", "Rationale for and methods of appraisal, if done."),
    ("METHODS", "13", "Synthesis of results", "Methods of handling and summarizing the charted data."),
    ("RESULTS", "14", "Selection of sources of evidence", "Numbers screened, assessed and included, with reasons for exclusion, ideally in a flow diagram."),
    ("RESULTS", "15", "Characteristics of sources of evidence", "Characteristics of each source for which data were charted, with citations."),
    ("RESULTS", "16", "Critical appraisal (optional)", "Data on critical appraisal of included sources, if done."),
    ("RESULTS", "17", "Results of individual sources", "Relevant data charted for each included source."),
    ("RESULTS", "18", "Synthesis of results", "Summary and/or presentation of the charting results related to the objectives."),
    ("DISCUSSION", "19", "Summary of evidence", "Main results, links to concepts and relevance to key groups."),
    ("DISCUSSION", "20", "Limitations", "Limitations of the scoping review process."),
    ("DISCUSSION", "21", "Conclusions", "General interpretation with respect to objectives, potential implications and next steps."),
    ("FUNDING", "22", "Funding", "Sources of funding and role of funders."),
]


def main():
    d = pd.read_csv(ROOT / "data" / "processed" / "included_studies_coded.csv")
    d["serotypes_involved"] = d["serotype"].apply(
        lambda s: ", ".join(f"DENV-{n}" for n in sorted(set(re.findall(r"DENV-?(\d)", s.replace("DENV-1 a DENV-4", "DENV-1 DENV-2 DENV-3 DENV-4"))))) or "-")
    s1 = d[["id", "author", "year", "title", "journal", "doi", "language_en", "databases", "states", "region",
            "scale", "design", "period_start", "period_end", "serotypes_involved", "serotype_status",
            "serotype", "setting", "main_findings"]].rename(columns={
        "id": "ID", "author": "First author", "year": "Year", "title": "Title", "journal": "Journal",
        "doi": "DOI", "language_en": "Language", "databases": "Source(s)", "states": "Federative unit(s)",
        "region": "Region", "scale": "Spatial scale", "design": "Design/focus", "period_start": "Epidemic period (start)",
        "period_end": "Epidemic period (end)", "serotypes_involved": "Serotypes involved",
        "serotype_status": "Source of serotype information",
        "serotype": "Serotype detail (original charting, Portuguese)", "setting": "Setting (original charting, Portuguese)",
        "main_findings": "Main findings (original charting, Portuguese) [INSERT: English translation]"})
    with pd.ExcelWriter(SUP / "S1_included_studies.xlsx") as w:
        s1.to_excel(w, sheet_name="S1 Included studies", index=False)
        pd.DataFrame({"Note": [
            "Serotype detail: † = extracted from full text; * = not reported, inferred from other included studies.",
            "Design/focus codes follow the codebook in code/01_build_dataset.py.",
            "Main findings were charted in Portuguese by the review team; [INSERT: English translation before submission]."]}
        ).to_excel(w, sheet_name="Notes", index=False)
    log = pd.read_csv(ROOT / "data" / "raw" / "search_log.csv")
    log.to_excel(SUP / "S2_search_log.xlsx", index=False)

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
    doc.add_paragraph().add_run("Supplementary Material S3. PRISMA-ScR checklist").bold = True
    t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
    for i, h in enumerate(["Section", "Item", "Checklist item", "Reported on page #"]):
        t.rows[0].cells[i].text = h
    for sec, n, item, desc in PRISMA_SCR:
        c = t.add_row().cells
        c[0].text, c[1].text, c[2].text = sec, n, f"{item}: {desc}"
        r = c[3].paragraphs[0].add_run("[INSERT]"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph("Source: Tricco AC et al. Ann Intern Med 2018; 169(7):467-473.")
    doc.save(SUP / "S3_PRISMA-ScR_checklist.docx")
    print("Supplementary files written")


if __name__ == "__main__":
    main()
